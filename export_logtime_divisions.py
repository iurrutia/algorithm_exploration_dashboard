#!/usr/bin/env python3
"""Create two color-coded Log-Time Mean ranking workbooks.

Place this file in the repository root beside ``build_dashboard.py``, then run:

    python build_dashboard.py
    python export_logtime_divisions_standalone.py

Outputs (by default under ``output/public_rankings``):

* ``logtime_mean_with_scores_and_divisions.xlsx``
* ``logtime_mean_divisions_only.xlsx``

The first workbook retains the Log-Time Volatility Mean score.  The second
keeps rank and division but omits the score everywhere, including the monthly
history sheets.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent
CODE_DIR = ROOT / "code"
sys.path.insert(0, str(CODE_DIR))

from jpar_pipeline import (  # type: ignore[import-not-found]  # noqa: E402
    make_event_time_records_df,
    make_qualified_events_df,
)


CALCULATION_CSV = ROOT / "data/data_jpar_v2/source_of_truth_calculation_df.csv"
JPAR_RESULTS_CSV = ROOT / "data/data_jpar_v2/source_of_truth_jpar_results.csv"
EXTERNAL_LEADERBOARD_CSV = ROOT / "data/colleague_systems/final_leaderboard.csv"
EXTERNAL_HISTORY_CSV = ROOT / "data/colleague_systems/state_history.csv"

HEADER_FILL = PatternFill("solid", fgColor="4F81BD")
HEADER_FONT = Font(color="FFFFFF", bold=True)
TITLE_FILL = PatternFill("solid", fgColor="D9EAF7")


@dataclass(frozen=True)
class RankingMethod:
    key: str
    label: str
    sheet_label: str
    filename: str
    higher_is_better: bool


METHOD = RankingMethod(
    "external_logtime",
    "Log-Time Volatility Mean",
    "Log-Time Mean",
    "logtime_volatility_mean_rankings.xlsx",
    True,
)


def member_key(value: object) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip()
    return text[:-2] if re.fullmatch(r"\d+\.0", text) else text


def event_key(value: object) -> str:
    if pd.isna(value):
        return ""
    text = str(value).strip()
    return text[:-2] if re.fullmatch(r"\d+\.0", text) else text


def ranked_frame(
    names: pd.DataFrame,
    score_col: str,
    event_col: str,
    higher_is_better: bool,
) -> pd.DataFrame:
    frame = names[["member_key", "full_name", score_col, event_col]].copy()
    frame = frame.rename(
        columns={score_col: "Rating", event_col: "# Events", "full_name": "Name"}
    )
    frame["Rating"] = pd.to_numeric(frame["Rating"], errors="coerce")
    frame["# Events"] = (
        pd.to_numeric(frame["# Events"], errors="coerce").fillna(0).astype(int)
    )
    frame = frame.dropna(subset=["Rating"])
    frame = frame[frame["Name"].fillna("").astype(str).str.strip().ne("")]
    frame = frame.sort_values(
        ["Rating", "Name", "member_key"],
        ascending=[not higher_is_better, True, True],
        kind="stable",
    ).reset_index(drop=True)
    frame.insert(
        0,
        "Rank",
        frame["Rating"].rank(method="min", ascending=not higher_is_better).astype(int),
    )
    return frame[["Rank", "Rating", "Name", "# Events", "member_key"]]


def participant_updates(history: pd.DataFrame, score_col: str) -> pd.DataFrame:
    h = history.copy()
    h["event_date"] = pd.to_datetime(h["event_date"], errors="coerce")
    h["event_id"] = h["event_id"].map(event_key)
    h["member_key"] = h["member_key"].map(member_key)
    h["events"] = pd.to_numeric(h["events"], errors="coerce").fillna(0).astype(int)
    h[score_col] = pd.to_numeric(h[score_col], errors="coerce")
    h = h.sort_values(["member_key", "event_date", "event_id"], kind="stable")
    previous = h.groupby("member_key")["events"].shift(fill_value=0)
    return h[h["events"].gt(previous)].copy()


def external_data(
    leaderboard: pd.DataFrame,
    history: pd.DataFrame,
    score_col: str,
) -> tuple[pd.DataFrame, pd.DataFrame, str, str]:
    latest = leaderboard.copy()
    latest["member_key"] = latest["member_key"].map(member_key)
    return latest, participant_updates(history, score_col), score_col, "external_events"


def monthly_frames(
    updates: pd.DataFrame,
    score_col: str,
    names: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    d = updates.dropna(subset=["event_date", score_col]).copy()
    d["Month"] = d["event_date"].dt.to_period("M")
    if d.empty:
        empty = pd.DataFrame(columns=["Month", "Name", "Rating", "Event #"])
        return empty, pd.DataFrame(columns=["Name"])

    monthly_score = (
        d.sort_values(["member_key", "event_date", "event_id"], kind="stable")
        .groupby(["member_key", "Month"], as_index=False)[score_col]
        .last()
    )
    monthly_events = (
        d.groupby(["member_key", "Month"], as_index=False)["event_id"]
        .nunique()
        .rename(columns={"event_id": "Event #"})
    )
    months = pd.period_range(d["Month"].min(), d["Month"].max(), freq="M")
    people = names[["member_key", "full_name"]].drop_duplicates("member_key")
    grid = (
        people.assign(_join=1)
        .merge(pd.DataFrame({"Month": months, "_join": 1}), on="_join")
        .drop(columns="_join")
    )
    grid = grid.merge(monthly_score, on=["member_key", "Month"], how="left")
    grid = grid.merge(monthly_events, on=["member_key", "Month"], how="left")
    grid = grid.sort_values(["member_key", "Month"], kind="stable")
    grid[score_col] = grid.groupby("member_key")[score_col].ffill()
    grid["Event #"] = grid["Event #"].fillna(0).astype(int)
    grid["Month"] = grid["Month"].astype(str)
    grid = grid.rename(columns={"full_name": "Name", score_col: "Rating"})
    monthly_list = grid[["Month", "Name", "Rating", "Event #"]].sort_values(
        ["Name", "Month"], kind="stable"
    )
    monthly_grid = monthly_list.pivot_table(
        index="Name", columns="Month", values="Rating", aggfunc="last", sort=False
    ).reset_index()
    return monthly_list.reset_index(drop=True), monthly_grid


def excel_value(value: object) -> object:
    if pd.isna(value):
        return None
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return float(value)
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def add_dataframe(ws, frame: pd.DataFrame, table_name: str | None = None) -> None:
    ws.append([str(column) for column in frame.columns])
    for row in frame.itertuples(index=False, name=None):
        ws.append([excel_value(value) for value in row])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if table_name and ws.max_row > 1:
        table = Table(displayName=table_name, ref=ws.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)


def style_sheet(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 22
    for column_cells in ws.iter_cols():
        letter = get_column_letter(column_cells[0].column)
        width = max(
            (len(str(cell.value)) for cell in column_cells[:250] if cell.value is not None),
            default=8,
        ) + 2
        ws.column_dimensions[letter].width = min(max(width, 10), 42)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top")


def add_qualified_events(ws, qualified: pd.DataFrame, method: RankingMethod) -> None:
    ws.append([f"{method.label} is calculated from the qualifying events listed below."])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(qualified.columns))
    ws["A1"].fill = TITLE_FILL
    ws["A1"].font = Font(bold=True, color="1F4E78")
    ws["A1"].alignment = Alignment(wrap_text=True)
    ws.append(list(qualified.columns))
    for row in qualified.itertuples(index=False, name=None):
        ws.append([excel_value(value) for value in row])
    for cell in ws[2]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"


def add_event_times(ws, event_times: pd.DataFrame) -> None:
    display = event_times.drop(columns=["Member ID"], errors="ignore").rename(
        columns={"Name": "Full Name"}
    )
    add_dataframe(ws, display)
    ws.freeze_panes = "B2"
    for column in range(2, ws.max_column + 1):
        for cells in ws.iter_cols(min_col=column, max_col=column, min_row=2):
            for cell in cells:
                cell.alignment = Alignment(horizontal="center")


def require_files(paths: list[Path]) -> None:
    missing = [path for path in paths if not path.exists()]
    if missing:
        formatted = "\n".join(f"  - {path.relative_to(ROOT)}" for path in missing)
        raise SystemExit(
            f"Missing build output(s):\n{formatted}\nRun `python build_dashboard.py` first."
        )


SCORE_COLUMN = "external_logtime"

# Valorant convention: division 3 is the highest division within a belt.
# Elite is deliberately a top-100 placement designation; all other divisions
# use stable score thresholds.
DIVISIONS = (
    "Elite",
    "Platinum 3",
    "Platinum 2",
    "Platinum 1",
    "Gold 3",
    "Gold 2",
    "Gold 1",
    "Silver 3",
    "Silver 2",
    "Silver 1",
    "Bronze 3",
    "Bronze 2",
    "Bronze 1",
)

DIVISION_COLORS = {
    "Elite": "7E57C2",
    "Platinum 3": "168C95",
    "Platinum 2": "42AEB5",
    "Platinum 1": "8FD3D7",
    "Gold 3": "D5A900",
    "Gold 2": "EACB3A",
    "Gold 1": "F7E58B",
    "Silver 3": "7D8B92",
    "Silver 2": "AAB5BA",
    "Silver 1": "D6DDE0",
    "Bronze 3": "8C542C",
    "Bronze 2": "B97849",
    "Bronze 1": "D7A47E",
}

DARK_TEXT_DIVISIONS = {
    "Platinum 1", "Gold 2", "Gold 1", "Silver 2", "Silver 1",
    "Bronze 2", "Bronze 1",
}


def division_for(score: float, rank: int | None = None) -> str:
    """Return the published division for one score/rank pair."""
    if rank is not None and rank <= 100:
        return "Elite"
    if score >= 0.30:
        return "Platinum 3"
    if score >= 0.20:
        return "Platinum 2"
    if score >= 0.10:
        return "Platinum 1"
    if score >= 0.00:
        return "Gold 3"
    if score >= -0.08:
        return "Gold 2"
    if score >= -0.17:
        return "Gold 1"
    if score >= -0.28:
        return "Silver 3"
    if score >= -0.39:
        return "Silver 2"
    if score >= -0.50:
        return "Silver 1"
    if score >= -0.65:
        return "Bronze 3"
    if score >= -0.85:
        return "Bronze 2"
    return "Bronze 1"


def add_current_divisions(ranked: pd.DataFrame) -> pd.DataFrame:
    out = ranked.copy()
    out["Division"] = [
        division_for(float(score), int(rank))
        for score, rank in zip(out["Rating"], out["Rank"], strict=True)
    ]
    return out


def add_monthly_divisions(monthly_list: pd.DataFrame) -> pd.DataFrame:
    """Assign Elite using each month's top 100, then apply score bands."""
    out = monthly_list.copy()
    out["Monthly Rank"] = out.groupby("Month")["Rating"].rank(
        method="min", ascending=False, na_option="bottom"
    )
    out["Division"] = [
        None if pd.isna(score) else division_for(float(score), int(rank))
        for score, rank in zip(out["Rating"], out["Monthly Rank"], strict=True)
    ]
    return out.drop(columns="Monthly Rank")


def color_division_cells(ws, division_header: str = "Division") -> None:
    headers = {cell.value: cell.column for cell in ws[1]}
    division_col = headers.get(division_header)
    rank_col = headers.get("Rank")
    if division_col is None:
        return

    for row in range(2, ws.max_row + 1):
        division = ws.cell(row, division_col).value
        color = DIVISION_COLORS.get(division)
        if color is None:
            continue
        font_color = "1F2937" if division in DARK_TEXT_DIVISIONS else "FFFFFF"
        for col in (division_col, rank_col):
            if col is None:
                continue
            cell = ws.cell(row, col)
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(bold=True, color=font_color)
            cell.alignment = Alignment(horizontal="center", vertical="center")


def add_distribution_sheet(wb: Workbook, ranked: pd.DataFrame) -> None:
    ws = wb.create_sheet("Division Distribution")
    counts = ranked["Division"].value_counts().reindex(DIVISIONS, fill_value=0)
    total = int(counts.sum())

    ws.append(["Division", "Puzzlers", "Share"])
    for division, count in counts.items():
        ws.append([division, int(count), int(count) / total if total else 0])

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor="263238")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:C{ws.max_row}"
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    ws["C2"].number_format = "0.0%"
    for row in range(2, ws.max_row + 1):
        ws.cell(row, 3).number_format = "0.0%"
        division = ws.cell(row, 1).value
        color = DIVISION_COLORS[division]
        font_color = "1F2937" if division in DARK_TEXT_DIVISIONS else "FFFFFF"
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=color)
        ws.cell(row, 1).font = Font(bold=True, color=font_color)

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "Puzzlers by Division"
    chart.x_axis.title = "Number of puzzlers"
    chart.y_axis.title = "Division"
    chart.height = 9
    chart.width = 17
    chart.legend = None
    chart.gapWidth = 45
    chart.dLbls = DataLabelList()
    chart.dLbls.showVal = True
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
    chart.series[0].dPt = [
        DataPoint(
            idx=index,
            spPr=GraphicalProperties(solidFill=DIVISION_COLORS[division]),
        )
        for index, division in enumerate(DIVISIONS)
    ]
    ws.add_chart(chart, "E2")
    ws.sheet_view.showGridLines = False


def ranking_views(ranked: pd.DataFrame, include_scores: bool) -> tuple[pd.DataFrame, pd.DataFrame]:
    columns = ["Rank"]
    if include_scores:
        columns.append("Rating")
    columns += ["Division", "Name", "# Events"]
    by_rank = ranked[columns].copy()
    by_name = by_rank.sort_values(["Name", "Rank"], kind="stable").reset_index(drop=True)
    if include_scores:
        by_rank = by_rank.rename(columns={"Rating": METHOD.label})
        by_name = by_name.rename(columns={"Rating": METHOD.label})
    return by_name, by_rank


def public_monthly_views(monthly: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    monthly_list = monthly[["Month", "Name", "Division", "Event #"]].copy()
    monthly_grid = monthly_list.pivot_table(
        index="Name", columns="Month", values="Division", aggfunc="last", sort=False
    ).reset_index()
    return monthly_list, monthly_grid


def build_division_workbook(
    ranked: pd.DataFrame,
    monthly: pd.DataFrame,
    qualified: pd.DataFrame,
    event_times: pd.DataFrame,
    destination: Path,
    include_scores: bool,
) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Qualified Events")
    add_qualified_events(ws, qualified, METHOD)

    ws = wb.create_sheet("Event Time Records")
    add_event_times(ws, event_times)

    by_name, by_rank = ranking_views(ranked, include_scores)
    ws = wb.create_sheet("Log-Time Mean by Name")
    add_dataframe(ws, by_name, "LogTimeByName")
    color_division_cells(ws)

    ws = wb.create_sheet("Log-Time Mean by Rank")
    add_dataframe(ws, by_rank, "LogTimeByRank")
    color_division_cells(ws)

    if include_scores:
        monthly_list = monthly[["Month", "Name", "Rating", "Division", "Event #"]].rename(
            columns={"Rating": METHOD.label}
        )
        monthly_grid = monthly.dropna(subset=["Rating"]).pivot_table(
            index="Name", columns="Month", values="Rating", aggfunc="last", sort=False
        ).reset_index()
    else:
        monthly_list, monthly_grid = public_monthly_views(monthly)

    ws = wb.create_sheet("Monthly Splits-List")
    add_dataframe(ws, monthly_list, "MonthlySplitsList")
    color_division_cells(ws)

    ws = wb.create_sheet("Monthly Splits-Grid")
    add_dataframe(ws, monthly_grid)
    if not include_scores:
        for column in range(2, ws.max_column + 1):
            for row in range(2, ws.max_row + 1):
                division = ws.cell(row, column).value
                color = DIVISION_COLORS.get(division)
                if color:
                    font_color = "1F2937" if division in DARK_TEXT_DIVISIONS else "FFFFFF"
                    ws.cell(row, column).fill = PatternFill("solid", fgColor=color)
                    ws.cell(row, column).font = Font(color=font_color)

    add_distribution_sheet(wb, ranked)

    for sheet in wb.worksheets:
        if sheet.title != "Division Distribution":
            style_sheet(sheet)
    if include_scores:
        for sheet_name in ("Log-Time Mean by Name", "Log-Time Mean by Rank", "Monthly Splits-List", "Monthly Splits-Grid"):
            for row in wb[sheet_name].iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, float):
                        cell.number_format = "0.0000"

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destination)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=ROOT / "output/public_rankings",
        help="Destination directory (default: output/public_rankings)",
    )
    args = parser.parse_args()

    required = [CALCULATION_CSV, JPAR_RESULTS_CSV, EXTERNAL_LEADERBOARD_CSV, EXTERNAL_HISTORY_CSV]
    require_files(required)

    calculation = pd.read_csv(CALCULATION_CSV, low_memory=False, dtype={"event_id": "string"})
    results = pd.read_csv(
        JPAR_RESULTS_CSV,
        low_memory=False,
        dtype={"event_id": "string", "resolved_member_id": "string"},
    )
    leaderboard = pd.read_csv(EXTERNAL_LEADERBOARD_CSV, low_memory=False, dtype={"member_key": "string"})
    history = pd.read_csv(
        EXTERNAL_HISTORY_CSV,
        low_memory=False,
        dtype={"event_id": "string", "member_key": "string"},
    )

    qualified = make_qualified_events_df(calculation)
    event_times = make_event_time_records_df(results)
    latest, updates, score_col, event_col = external_data(leaderboard, history, SCORE_COLUMN)
    ranked = add_current_divisions(ranked_frame(latest, score_col, event_col, True))
    names = ranked[["member_key", "Name"]].rename(columns={"Name": "full_name"})
    monthly_list, _ = monthly_frames(updates, score_col, names)
    monthly = add_monthly_divisions(monthly_list)

    outputs = (
        ("logtime_mean_with_scores_and_divisions.xlsx", True),
        ("logtime_mean_divisions_only.xlsx", False),
    )
    written: list[Path] = []
    for filename, include_scores in outputs:
        destination = args.output_dir.resolve() / filename
        build_division_workbook(
            ranked, monthly, qualified, event_times, destination, include_scores
        )
        written.append(destination)

    print("Created:")
    for path in written:
        print(f"  {path}")


if __name__ == "__main__":
    main()
